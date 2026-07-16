import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

INPUT_PATH = Path(sys.argv[1])
OUTPUT_PATH = Path(sys.argv[2])

with INPUT_PATH.open("r", encoding="utf-8") as handle:
    payload = json.load(handle)

organization_name = str(payload.get("organizationName") or "University of Cebu - South Campus")
campus_address = str(payload.get("campusAddress") or "")
report_title = str(payload.get("reportTitle") or "ATTENDANCE REPORT")
classification = str(payload.get("classification") or "NON-TEACHING")
timeframe_label = str(payload.get("timeframeLabel") or "All attendance records")
rows = payload.get("rows") or []

wb = Workbook(write_only=True)
ws = wb.create_sheet(title="Attendance Report")

thin_side = Side(style="thin", color="D1D5DB")
header_fill = PatternFill(fill_type="solid", fgColor="E8EEF8")
meta_fill = PatternFill(fill_type="solid", fgColor="F8FAFC")
center = Alignment(horizontal="center", vertical="center")
meta_center = Alignment(horizontal="centerContinuous", vertical="center")
left = Alignment(horizontal="left", vertical="center")

column_widths = {
    "A": 28,
    "B": 24,
    "C": 16,
    "D": 10,
    "E": 12,
    "F": 12,
    "G": 12,
    "H": 12,
    "I": 14,
    "J": 10,
    "K": 10,
    "L": 10,
    "M": 18,
}

for column, width in column_widths.items():
    ws.column_dimensions[column].width = width


def make_cell(value, *, bold=False, fill=None, alignment=left):
    cell = WriteOnlyCell(ws, value=value)
    cell.font = Font(size=11, bold=bold)
    cell.alignment = alignment
    cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    if fill is not None:
        cell.fill = fill
    return cell


for text, font_size, bold in [
    (organization_name.upper(), 13, True),
    (campus_address.upper(), 11, False),
    (report_title.upper(), 12, True),
    (classification.upper(), 11, True),
    (timeframe_label, 11, True),
]:
    row = [make_cell(text, bold=bold, fill=meta_fill, alignment=meta_center)]
    row.extend(make_cell("", fill=meta_fill, alignment=meta_center) for _ in range(12))
    row[0].font = Font(size=font_size, bold=bold)
    ws.append(row)

ws.append([])

headers = [
    "Employee",
    "Position",
    "Date",
    "Day",
    "AM In",
    "AM Out",
    "PM In",
    "PM Out",
    "Work Hours",
    "Late",
    "UT",
    "OT",
    "Remarks",
]
ws.append([make_cell(header, bold=True, fill=header_fill, alignment=center) for header in headers])

for record in rows:
    ws.append([
        make_cell(record.get("employeeName") or ""),
        make_cell(record.get("position") or ""),
        make_cell(record.get("dateLabel") or "", alignment=center),
        make_cell(record.get("day") or "", alignment=center),
        make_cell(record.get("morningIn") or "-", alignment=center),
        make_cell(record.get("morningOut") or "-", alignment=center),
        make_cell(record.get("afternoonIn") or "-", alignment=center),
        make_cell(record.get("afternoonOut") or "-", alignment=center),
        make_cell(record.get("workHours") or "-", alignment=center),
        make_cell(record.get("late") or "-", alignment=center),
        make_cell(record.get("undertime") or "-", alignment=center),
        make_cell(record.get("overtime") or "-", alignment=center),
        make_cell(record.get("remarks") or "", alignment=center),
    ])

OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT_PATH)

