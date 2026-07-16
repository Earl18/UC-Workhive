import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

INPUT_PATH = Path(sys.argv[1])
OUTPUT_PATH = Path(sys.argv[2])

with INPUT_PATH.open("r", encoding="utf-8") as handle:
    payload = json.load(handle)

organization_name = str(payload.get("organizationName") or "University of Cebu - South Campus")
campus_address = str(payload.get("campusAddress") or "")
report_title = str(payload.get("reportTitle") or "ATTENDANCE SUMMARY")
classification = str(payload.get("classification") or "NON-TEACHING PERSONNEL")
timeframe_label = str(payload.get("timeframeLabel") or "")
period_label = str(payload.get("periodLabel") or "")
rows = payload.get("rows") or []

wb = Workbook()
ws = wb.active
ws.title = "Attendance Summary"

thin_side = Side(style="thin", color="000000")
medium_side = Side(style="medium", color="000000")
all_thin = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
all_medium = Border(left=medium_side, right=medium_side, top=medium_side, bottom=medium_side)

title_fill = PatternFill(fill_type="solid", fgColor="9DC3E6")
header_fill = PatternFill(fill_type="solid", fgColor="BDD7EE")
section_fill = PatternFill(fill_type="solid", fgColor="C6E0B4")
white_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)
remarks_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

column_widths = {
    "A": 6,
    "B": 20,
    "C": 20,
    "D": 12,
    "E": 14,
    "F": 14,
    "G": 16,
    "H": 18,
    "I": 56,
}
for column, width in column_widths.items():
    ws.column_dimensions[column].width = width


def style_cell(cell, *, fill=None, font=None, alignment=None, border=None):
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border


def style_range(cell_range, *, fill=None, font=None, alignment=None, border=None):
    for row in ws[cell_range]:
        for cell in row:
            style_cell(cell, fill=fill, font=font, alignment=alignment, border=border)


def merge_label(cell_range, value, *, fill, font, alignment=center, border=all_thin):
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = value
    style_range(cell_range, fill=fill, font=font, alignment=alignment, border=border)


def format_remarks(value):
    text = str(value or "-").strip()
    if not text or text == "-":
        return "-"
    phrases = [part.strip() for part in text.split(",") if part.strip()]
    return ", ".join(phrases)


def estimate_row_height(*, remarks_text):
    estimated_lines = max(1, (len(remarks_text) // 54) + 1)
    return max(22, 18 * estimated_lines)


ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 20
ws.row_dimensions[3].height = 28
ws.row_dimensions[4].height = 22
ws.row_dimensions[5].height = 24
ws.row_dimensions[6].height = 36
ws.row_dimensions[7].height = 24

merge_label("A1:I1", organization_name.upper(), fill=white_fill, font=Font(size=13, bold=True), border=all_medium)
merge_label("A2:I2", campus_address.upper() if campus_address else "", fill=white_fill, font=Font(size=10), border=all_medium)
merge_label("A3:I3", report_title.upper(), fill=title_fill, font=Font(size=16, bold=True), border=all_medium)
merge_label("A4:I4", timeframe_label.upper() if timeframe_label else "", fill=title_fill, font=Font(size=12, bold=True), border=all_medium)
merge_label("A5:I5", period_label, fill=header_fill, font=Font(size=11, bold=True), border=all_medium)

ws.merge_cells("A6:A7")
ws["A6"] = "#"
style_range("A6:A7", fill=header_fill, font=Font(size=10, bold=True), alignment=center, border=all_medium)

ws.merge_cells("B6:C6")
ws["B6"] = "NAME OF EMPLOYEE"
style_range("B6:C6", fill=header_fill, font=Font(size=10, bold=True), alignment=center, border=all_medium)

ws["B7"] = "LAST NAME"
ws["C7"] = "FIRST NAME"
style_range("B7:C7", fill=header_fill, font=Font(size=10, bold=True), alignment=center, border=all_medium)

header_specs = [
    ("D6:D7", "ABSENT"),
    ("E6:E7", "LEAVE\nDAYS"),
    ("F6:F7", "LATE\n(minutes)"),
    ("G6:G7", "UNDERTIME\n(minutes)"),
    ("H6:H7", "OVER-TIME\n(Regular) hrs"),
    ("I6:I7", "REMARKS"),
]
for cell_range, label in header_specs:
    ws.merge_cells(cell_range)
    anchor = ws[cell_range.split(":")[0]]
    anchor.value = label
    style_range(cell_range, fill=header_fill, font=Font(size=10, bold=True), alignment=center, border=all_medium)

merge_label("A8:I8", classification.upper(), fill=section_fill, font=Font(size=11, bold=True), border=all_medium)

start_row = 9
for index, record in enumerate(rows, start=start_row):
    remarks_text = format_remarks(record.get("remarks"))
    ws.row_dimensions[index].height = estimate_row_height(remarks_text=remarks_text)
    ws.cell(index, 1, record.get("sequence") or index - start_row + 1)
    ws.cell(index, 2, record.get("lastName") or "")
    ws.cell(index, 3, record.get("firstName") or "")
    ws.cell(index, 4, record.get("absent") or "-")
    ws.cell(index, 5, record.get("leaveDays") or "-")
    ws.cell(index, 6, record.get("late") or "-")
    ws.cell(index, 7, record.get("undertime") or "-")
    ws.cell(index, 8, record.get("overtime") or "-")
    ws.cell(index, 9, remarks_text)

    for col in range(1, 10):
        cell = ws.cell(index, col)
        style_cell(
            cell,
            border=all_thin,
            alignment=remarks_alignment if col == 9 else left if col in (2, 3) else center,
            font=Font(size=10),
        )

ws.freeze_panes = "A9"

OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT_PATH)
