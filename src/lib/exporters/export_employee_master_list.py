import json
import sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INPUT_PATH = Path(sys.argv[1])
OUTPUT_PATH = Path(sys.argv[2])

with INPUT_PATH.open('r', encoding='utf-8') as handle:
    payload = json.load(handle)

organization_name = str(payload.get('organizationName') or 'University of Cebu - South Campus')
campus_address = str(payload.get('campusAddress') or '')
report_title = str(payload.get('reportTitle') or 'MONITORING REPORT OF EMPLOYMENT STATUS')
classification = str(payload.get('classification') or 'ADMINISTRATIVE / NON-TEACHING')
employees = payload.get('employees') or []

wb = Workbook()
ws = wb.active
ws.title = 'Employment Status'

last_column = 18
last_column_letter = get_column_letter(last_column)
thin_side = Side(style='thin', color='D1D5DB')
header_fill = PatternFill(fill_type='solid', fgColor='E8EEF8')
subheader_fill = PatternFill(fill_type='solid', fgColor='F4F7FB')
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)

for row_index, text, font_size, bold in [
    (1, organization_name.upper(), 13, True),
    (2, campus_address.upper(), 11, False),
    (3, report_title.upper(), 12, True),
    (4, classification.upper(), 11, True),
]:
    ws.merge_cells(f'A{row_index}:{last_column_letter}{row_index}')
    cell = ws[f'A{row_index}']
    cell.value = text
    cell.font = Font(size=font_size, bold=bold)
    cell.alignment = center

ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 20
ws.row_dimensions[3].height = 22
ws.row_dimensions[4].height = 20
ws.row_dimensions[6].height = 24
ws.row_dimensions[7].height = 22

header_merges = [
    ('A6:A7', 'ID NUMBER'),
    ('B6:D6', 'NAME'),
    ('E6:E7', 'GENDER'),
    ('F6:F7', 'DEPARTMENT'),
    ('G6:G7', 'POSITION'),
    ('H6:H7', 'DATE HIRED'),
    ('I6:I7', 'LENGTH OF SERVICE'),
    ('J6:J7', 'AGE'),
    ('K6:K7', 'E-STATUS'),
    ('L6:L7', 'EMAIL ADDRESS'),
    ('M6:O6', 'EDUCATIONAL ATTAINMENT'),
    ('P6:P7', 'ADDRESS'),
    ('Q6:Q7', 'CONTACT NUMBER'),
    ('R6:R7', 'BIRTHDATE'),
]

for merged_range, value in header_merges:
    ws.merge_cells(merged_range)
    cell = ws[merged_range.split(':')[0]]
    cell.value = value
    cell.font = Font(bold=True)
    cell.fill = header_fill
    cell.alignment = center

for column, value in [('B7', 'LAST NAME'), ('C7', 'FIRST NAME'), ('D7', 'MIDDLE NAME'), ('M7', 'BACHELOR'), ('N7', 'MASTER'), ('O7', 'DOCTORAL')]:
    cell = ws[column]
    cell.value = value
    cell.font = Font(bold=True)
    cell.fill = subheader_fill
    cell.alignment = center

column_widths = {
    'A': 6, 'B': 18, 'C': 18, 'D': 18, 'E': 12, 'F': 20, 'G': 26, 'H': 14,
    'I': 22, 'J': 8, 'K': 14, 'L': 28, 'M': 34, 'N': 34, 'O': 34, 'P': 28,
    'Q': 18, 'R': 14,
}
for column, width in column_widths.items():
    ws.column_dimensions[column].width = width

for row_number, employee in enumerate(employees, start=8):
    education = employee.get('education') or []
    edu_map = {str(item.get('level') or ''): str(item.get('degree') or '') for item in education if isinstance(item, dict)}
    values = [
        employee.get('idNumber') or '',
        employee.get('lastName') or '',
        employee.get('firstName') or '',
        employee.get('middleName') or '',
        employee.get('gender') or '',
        employee.get('department') or '',
        employee.get('position') or '',
        employee.get('dateHiredDisplay') or employee.get('dateHired') or '',
        employee.get('lengthOfService') or '',
        employee.get('age') if employee.get('age') is not None else '',
        employee.get('employmentStatus') or '',
        employee.get('email') or '',
        edu_map.get('Bachelor', ''),
        edu_map.get('Master', ''),
        edu_map.get('Doctoral', ''),
        employee.get('address') or '',
        employee.get('contactNumber') or '',
        employee.get('birthDate') or '',
    ]
    for column_index, value in enumerate(values, start=1):
        cell = ws.cell(row=row_number, column=column_index, value=value)
        cell.alignment = center if column_index in {1, 5, 8, 9, 10, 11, 17, 18} else left

for row in ws.iter_rows(min_row=6, max_row=max(7, 7 + len(employees)), min_col=1, max_col=last_column):
    for cell in row:
        cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

ws.freeze_panes = 'A8'
ws.sheet_view.showGridLines = True

OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT_PATH)
