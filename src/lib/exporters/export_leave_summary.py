import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

INPUT_PATH = Path(sys.argv[1])
OUTPUT_PATH = Path(sys.argv[2])

with INPUT_PATH.open("r", encoding="utf-8") as handle:
    payload = json.load(handle)

organization_name = str(payload.get("organizationName") or "UNIVERSITY OF CEBU - PARDO AND TALISAY")
report_title = str(payload.get("reportTitle") or "LEAVE SUMMARY REMAINING CREDITS")
timeframe_label = str(payload.get("timeframeLabel") or "LEAVE CREDITS FROM JANUARY - DECEMBER")
filters_label = str(payload.get("filtersLabel") or "")
rows = payload.get("rows") or []

wb = Workbook()
ws = wb.active
ws.title = "Leave Summary"

thin = Side(style="thin", color="000000")
medium = Side(style="medium", color="000000")
all_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
all_medium = Border(left=medium, right=medium, top=medium, bottom=medium)

blue_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
header_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
green_fill = PatternFill(fill_type="solid", fgColor="C6E0B4")
yellow_fill = PatternFill(fill_type="solid", fgColor="FFE699")
orange_fill = PatternFill(fill_type="solid", fgColor="F4B183")
white_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")

center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

widths = {
    "A": 6,
    "B": 18,
    "C": 18,
    "D": 18,
    "E": 26,
    "F": 14,
    "G": 16,
    "H": 14,
    "I": 14,
    "J": 16,
    "K": 14,
    "L": 16,
    "M": 16,
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width


def style_range(cell_range, *, fill=None, font=None, alignment=None, border=None):
    for row in ws[cell_range]:
        for cell in row:
            if fill is not None:
                cell.fill = fill
            if font is not None:
                cell.font = font
            if alignment is not None:
                cell.alignment = alignment
            if border is not None:
                cell.border = border


def merge_label(cell_range, value, *, fill, font, alignment=center, border=all_medium):
    ws.merge_cells(cell_range)
    anchor = ws[cell_range.split(":")[0]]
    anchor.value = value
    style_range(cell_range, fill=fill, font=font, alignment=alignment, border=border)

ws.row_dimensions[1].height = 28
ws.row_dimensions[2].height = 22
ws.row_dimensions[3].height = 20
ws.row_dimensions[4].height = 30

merge_label("A1:M1", organization_name.upper(), fill=blue_fill, font=Font(size=18, bold=True))
merge_label("A2:M2", report_title.upper(), fill=white_fill, font=Font(size=12, bold=True))
merge_label("A3:M3", timeframe_label.upper(), fill=white_fill, font=Font(size=11, bold=True))
merge_label("A4:M4", filters_label, fill=white_fill, font=Font(size=10, italic=bool(filters_label)))

headers = [
    ("A5", "#", header_fill),
    ("B5", "LAST NAME", header_fill),
    ("C5", "FIRST NAME", header_fill),
    ("D5", "MIDDLE NAME", header_fill),
    ("E5", "POSITION", header_fill),
    ("F5", "E-STATUS", header_fill),
    ("G5", "DATE HIRED", header_fill),
    ("H5", "VACATION LEAVE", green_fill),
    ("I5", "SICK LEAVE", yellow_fill),
    ("J5", "SUMMER INCENTIVE LEAVE", orange_fill),
    ("K5", "SOLO PARENT", green_fill),
    ("L5", "PATERNITY LEAVE", blue_fill),
    ("M5", "MATERNITY LEAVE", blue_fill),
]
for ref, value, fill in headers:
    ws[ref] = value
    style_range(f"{ref}:{ref}", fill=fill, font=Font(size=9, bold=True), alignment=center, border=all_medium)

start_row = 6
for row_index, record in enumerate(rows, start=start_row):
    ws.row_dimensions[row_index].height = 22
    values = [
        record.get("sequence") or row_index - start_row + 1,
        record.get("lastName") or "",
        record.get("firstName") or "",
        record.get("middleName") or "",
        record.get("position") or "",
        record.get("employmentStatus") or "",
        record.get("dateHired") or "-",
        record.get("vacationLeave") or "-",
        record.get("sickLeave") or "-",
        record.get("incentiveLeave") or "-",
        record.get("soloParent") or "-",
        record.get("paternity") or "-",
        record.get("maternity") or "-",
    ]
    for col_index, value in enumerate(values, start=1):
        cell = ws.cell(row_index, col_index, value)
        cell.border = all_thin
        cell.font = Font(size=10)
        cell.alignment = left if col_index in (2,3,4,5) else center

ws.freeze_panes = "A6"
OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUTPUT_PATH)
