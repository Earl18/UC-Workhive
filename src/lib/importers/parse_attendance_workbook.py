import json
import re
import sys
import datetime as dt

import openpyxl
from openpyxl.utils.datetime import from_excel

DATE_FORMATS = [
    "%Y-%m-%d",
    "%m/%d/%Y",
    "%m/%d/%y",
    "%B %d, %Y",
    "%b %d, %Y",
]
TIME_FORMATS = [
    "%I:%M %p",
    "%I:%M:%S %p",
    "%H:%M",
    "%H:%M:%S",
]


def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            if isinstance(converted, dt.datetime):
                return converted.date()
            if isinstance(converted, dt.date):
                return converted
        except Exception:
            return None
    text = normalize_text(value)
    for fmt in DATE_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def normalize_time(value):
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.strftime("%H:%M")
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            if isinstance(converted, dt.datetime):
                return converted.strftime("%H:%M")
            if isinstance(converted, dt.time):
                return converted.strftime("%H:%M")
        except Exception:
            return None
    text = normalize_text(value)
    if not text:
        return None
    for fmt in TIME_FORMATS:
        try:
            return dt.datetime.strptime(text, fmt).strftime("%H:%M")
        except ValueError:
            continue
    match = re.search(r"(\d{1,2}):(\d{2})", text)
    if match:
        hour = int(match.group(1))
        minute = int(match.group(2))
        if "PM" in text.upper() and hour < 12:
            hour += 12
        if "AM" in text.upper() and hour == 12:
            hour = 0
        if 0 <= hour <= 23 and 0 <= minute <= 59:
            return f"{hour:02d}:{minute:02d}"
    return None


def find_header_row(worksheet, max_row=12):
    for row_index, row in enumerate(worksheet.iter_rows(min_row=1, max_row=max_row, values_only=True), start=1):
        values = [normalize_text(cell).lower() for cell in row]
        if "employee name" in values and "date" in values and values.count("in") >= 2 and values.count("out") >= 2:
            return row_index, values
    raise ValueError("Unable to locate the Daily Time Record header row.")


def build_column_map(headers):
    in_columns = [index for index, value in enumerate(headers) if value == "in"]
    out_columns = [index for index, value in enumerate(headers) if value == "out"]
    return {
        "employee_name": headers.index("employee name"),
        "date": headers.index("date"),
        "day": headers.index("day"),
        "morning_in": in_columns[0],
        "morning_out": out_columns[0],
        "afternoon_in": in_columns[1],
        "afternoon_out": out_columns[1],
        "next_day": headers.index("next day") if "next day" in headers else None,
        "hours_work": headers.index("hours work") if "hours work" in headers else None,
    }


path = sys.argv[1]
workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
worksheet = workbook["Daily Time Record"] if "Daily Time Record" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
header_row, headers = find_header_row(worksheet)
column_map = build_column_map(headers)
rows = []
current_employee_name = ""

for row_index, row in enumerate(worksheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
    employee_name = normalize_text(row[column_map["employee_name"]]) if column_map["employee_name"] < len(row) else ""
    if employee_name:
        current_employee_name = employee_name
    work_date = normalize_date(row[column_map["date"]] if column_map["date"] < len(row) else None)
    if not work_date:
        continue
    if not current_employee_name:
        continue

    day_name = normalize_text(row[column_map["day"]] if column_map["day"] < len(row) else "") or work_date.strftime("%A")
    next_day_value = normalize_time(row[column_map["next_day"]]) if column_map["next_day"] is not None and column_map["next_day"] < len(row) else None
    hours_work = normalize_text(row[column_map["hours_work"]]) if column_map["hours_work"] is not None and column_map["hours_work"] < len(row) else ""

    rows.append({
        "rowNumber": row_index,
        "employeeName": current_employee_name,
        "date": work_date.isoformat(),
        "dayName": day_name,
        "morningIn": normalize_time(row[column_map["morning_in"]] if column_map["morning_in"] < len(row) else None),
        "morningOut": normalize_time(row[column_map["morning_out"]] if column_map["morning_out"] < len(row) else None),
        "afternoonIn": normalize_time(row[column_map["afternoon_in"]] if column_map["afternoon_in"] < len(row) else None),
        "afternoonOut": normalize_time(row[column_map["afternoon_out"]] if column_map["afternoon_out"] < len(row) else None),
        "nextDay": next_day_value,
        "rawHoursWork": hours_work,
    })

payload = {
    "sheetName": worksheet.title,
    "rows": rows,
}
print(json.dumps(payload, ensure_ascii=False))
