import json
import re
import sys
import datetime as dt
import unicodedata
import openpyxl
from openpyxl.utils.datetime import from_excel


def repair_mojibake(text):
    if not text:
        return text
    suspects = ('?', '?', '?', '?', '?')
    if any(token in text for token in suspects):
        try:
            repaired = text.encode('latin1', errors='ignore').decode('utf-8', errors='ignore').strip()
            if repaired:
                return repaired
        except Exception:
            pass
    return text


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    text = repair_mojibake(text)
    return unicodedata.normalize('NFC', text)


DATE_FORMATS = (
    '%Y-%m-%d', '%Y/%m/%d',
    '%m/%d/%Y', '%m/%d/%y', '%m-%d-%Y', '%m-%d-%y',
    '%m.%d.%Y', '%m.%d.%y', '%d-%b-%Y', '%d-%b-%y',
    '%b %d, %Y', '%b %d,%Y', '%B %d, %Y', '%B %d,%Y',
    '%b %d %Y', '%B %d %Y',
)

DATE_PATTERN_ALPHA = re.compile(r'[A-Za-z]')
DATE_PATTERN_ISO = re.compile(r'^(\d{4})[-/](\d{1,2})[-/](\d{1,2})$')
DATE_PATTERN_SLASH_DASH = re.compile(r'^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})$')
DATE_TEXT_CACHE = {}
DATE_FORMAT_LOOKUP = {
    '-': ('%Y-%m-%d', '%m-%d-%Y', '%m-%d-%y', '%d-%b-%Y', '%d-%b-%y'),
    '/': ('%Y/%m/%d', '%m/%d/%Y', '%m/%d/%y'),
    '.': ('%m.%d.%Y', '%m.%d.%y'),
    ',': ('%b %d, %Y', '%b %d,%Y', '%B %d, %Y', '%B %d,%Y'),
    ' ': ('%b %d %Y', '%B %d %Y'),
}


def parse_date_value(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            excel_date = from_excel(value)
            if isinstance(excel_date, dt.datetime):
                return excel_date.date()
            if isinstance(excel_date, dt.date):
                return excel_date
        except Exception:
            pass
    text = clean(value)
    if not text:
        return None

    if text in DATE_TEXT_CACHE:
        return DATE_TEXT_CACHE[text]

    compacted = re.sub(r'\s+', ' ', text)
    match = DATE_PATTERN_ISO.match(compacted)
    if match:
        year, month, day = (int(part) for part in match.groups())
        try:
            result = dt.date(year, month, day)
        except ValueError:
            result = None
        DATE_TEXT_CACHE[text] = result
        return result

    match = DATE_PATTERN_SLASH_DASH.match(compacted)
    if match:
        month, day, raw_year = match.groups()
        year = int(raw_year)
        if len(raw_year) == 2:
            year += 2000
        try:
            result = dt.date(year, int(month), int(day))
        except ValueError:
            result = None
        DATE_TEXT_CACHE[text] = result
        return result

    separator = ',' if ',' in compacted else '/' if '/' in compacted else '-' if '-' in compacted else '.' if '.' in compacted else ' ' if ' ' in compacted else ''
    formats = DATE_FORMAT_LOOKUP.get(separator, ())
    if not formats:
        if DATE_PATTERN_ALPHA.search(compacted):
            formats = DATE_FORMAT_LOOKUP[' '] + DATE_FORMAT_LOOKUP[','] + ('%d-%b-%Y', '%d-%b-%y')
        else:
            formats = DATE_FORMATS

    for fmt in formats:
        try:
            result = dt.datetime.strptime(compacted, fmt).date()
            DATE_TEXT_CACHE[text] = result
            return result
        except ValueError:
            pass

    DATE_TEXT_CACHE[text] = None
    return None


def normalize_date(value):
    parsed = parse_date_value(value)
    if parsed:
        return parsed.strftime('%Y-%m-%d')
    text = clean(value)
    return text[:10] if text else ''


def normalize_birth_date(value):
    parsed = parse_date_value(value)
    if not parsed:
        return ''

    today = dt.date.today()
    age = today.year - parsed.year - ((today.month, today.day) < (parsed.month, parsed.day))
    if age < 15 or age > 100:
        return ''

    return f"{parsed.month}/{parsed.day}/{parsed.year}"


def normalize_contact_number(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(int(value))
    digits = re.sub(r'\D+', '', clean(value))
    return digits


def normalize_gender(value):
    text = clean(value).lower()
    return 'Male' if text.startswith('m') else 'Female'


def normalize_status(value):
    text = clean(value).lower()
    if text == 'probationary':
        return 'Probationary'
    if text == 'retired':
        return 'Retired'
    if text == 'resigned':
        return 'Resigned'
    return 'Regular'


def strip_parenthetical(value):
    return re.sub(r'\s*\([^)]*\)', '', clean(value)).strip()


def normalize_name(value):
    text = clean(value)
    return text.title() if text else ''


def normalize_department(value):
    text = clean(value)
    if not text:
        return ''
    if text.upper() == text and len(text) <= 4:
        return text
    return text.title()


def normalize_header(value):
    text = clean(value).lower()
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def has_header_keywords(value, required_groups):
    normalized = f" {normalize_header(value)} "
    return all(any(f" {keyword} " in normalized for keyword in group) for group in required_groups)


def find_column_index(header_values, aliases=(), keyword_groups=()):
    normalized_aliases = {normalize_header(alias) for alias in aliases}
    for index, value in enumerate(header_values):
        normalized = normalize_header(value)
        if normalized in normalized_aliases:
            return index
        if keyword_groups and has_header_keywords(normalized, keyword_groups):
            return index
    return None


def find_header_row(worksheet, max_row=12):
    expected_aliases = {
        'id number', 'id no', 'id #', 'new id number', 'employee id', 'employee id number', 'sequence', 'no', '#', 'last name', 'firstname', 'first name', 'middle name', 'gender',
        'department', 'position', 'date hired', 'employment status', 'email', 'address',
        'date of birth', 'birth date', 'birthday', 'birth day', 'dob', 'contact number',
        'contact #', 'contact no', 'cellphone number', 'mobile number'
    }
    best_row_index = 7
    best_header_values = ()
    best_score = -1
    for row_index, row in enumerate(worksheet.iter_rows(min_row=1, max_row=max_row, values_only=True), start=1):
        header_values = tuple(row)
        normalized_values = [normalize_header(value) for value in header_values if normalize_header(value)]
        density_score = len(normalized_values)
        header_score = sum(1 for value in normalized_values if value in expected_aliases)
        keyword_score = sum(
            1 for value in normalized_values
            if any(token in value.split() for token in ('name', 'gender', 'department', 'position', 'date', 'email', 'address', 'contact', 'birth'))
        )
        score = (header_score * 10) + (keyword_score * 3) + density_score
        if score > best_score:
            best_score = score
            best_row_index = row_index
            best_header_values = header_values
    return best_row_index, best_header_values


def is_header_like_row(row):
    first_cells = [normalize_header(value) for value in row[:15] if normalize_header(value)]
    if not first_cells:
        return False

    explicit_headers = {
        'id number', 'id no', 'id #', 'new id number', 'employee id', 'employee id number', 'sequence', 'no', '#', 'last name', 'firstname', 'first name', 'middle name', 'gender',
        'department', 'position', 'date hired', 'employment status', 'email', 'address',
        'date of birth', 'birth date', 'birthday', 'birth day', 'dob', 'contact number',
        'contact #', 'contact no', 'cellphone number', 'mobile number'
    }
    matches = sum(1 for value in first_cells if value in explicit_headers)
    keyword_hits = sum(
        1 for value in first_cells
        if any(token in value.split() for token in ('name', 'gender', 'department', 'position', 'date', 'email', 'address', 'contact', 'birth', 'status'))
    )
    return matches >= 2 or keyword_hits >= 4




def detect_workbook_type(header_row, indexes):
    normalized_headers = [normalize_header(value) for value in header_row if normalize_header(value)]
    header_set = set(normalized_headers)

    has_explicit_id_header = any(header in header_set for header in (
        'new id number',
        'id number',
        'employee id number',
        'employee id',
        'id no',
        'id #',
    ))
    only_identity_columns = (
        indexes['id_number'] is not None
        and indexes['last_name'] is not None
        and indexes['first_name'] is not None
        and indexes['middle_name'] is not None
        and indexes['email'] is None
        and indexes['department'] is None
        and indexes['position'] is None
        and indexes['date_hired'] is None
        and indexes['status'] is None
    )

    if has_explicit_id_header and only_identity_columns:
        return 'id-number', 'ID Number Update File', 'Detected from headers like NEW ID NUMBER, LAST NAME, FIRST NAME, and MIDDLE NAME.'

    if has_explicit_id_header and indexes['email'] is None and indexes['department'] is None and indexes['position'] is None and indexes['date_hired'] is None:
        return 'id-number', 'ID Number Update File', 'Detected from an ID number-focused sheet with employee name columns.'

    return 'employee-information', 'Employee Masterlist File', 'Detected from employee information fields such as department, position, date hired, email, or employment status.'

def parse_education(level, degree):
    value = clean(degree)
    if not value:
        return None
    return {'level': level, 'degree': value, 'school': 'University of Cebu - South Campus'}


path = sys.argv[1]
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb['Employment Status'] if 'Employment Status' in wb.sheetnames else wb[wb.sheetnames[0]]
classification = clean(ws['A4'].value) or 'ADMINISTRATIVE / NON-TEACHING'
header_row_index, header_row = find_header_row(ws)
birth_date_index = find_column_index(
    header_row,
    aliases=['date of birth', 'birth date', 'birthday', 'birth day', 'dob', 'bday', 'birthdate', 'employee birthday', 'employee birth date']
)
contact_number_index = find_column_index(
    header_row,
    aliases=['contact number', 'contact #', 'contact#', 'contact no', 'contact no.', 'cellphone number', 'mobile number', 'phone number', 'phone no', 'telephone number', 'tel no', 'tel #', 'cp no', 'cp #', 'mobile no', 'employee contact number'],
    keyword_groups=[('contact', 'phone', 'telephone', 'mobile', 'cellphone', 'cp', 'tel'), ('number', 'no', '#')]
)

id_number_index = find_column_index(
    header_row,
    aliases=['id number', 'id no', 'id #', 'new id number', 'employee id', 'employee id number', 'employee no', 'idnum', 'id']
)
last_name_index = find_column_index(header_row, aliases=['last name'])
first_name_index = find_column_index(header_row, aliases=['firstname', 'first name'])
middle_name_index = find_column_index(header_row, aliases=['middle name', 'middlename', 'middle initial'])
gender_index = find_column_index(header_row, aliases=['gender'])
department_index = find_column_index(header_row, aliases=['department'])
position_index = find_column_index(header_row, aliases=['position'])
date_hired_index = find_column_index(header_row, aliases=['date hired', 'date of hire', 'hired date'])
status_index = find_column_index(header_row, aliases=['employment status', 'status', 'e-status', 'e status'])
email_index = find_column_index(header_row, aliases=['email', 'email address', 'uc email'])
address_index = find_column_index(header_row, aliases=['address', 'home address'])
bachelor_index = find_column_index(header_row, aliases=['bachelor', 'bachelors', 'bachelor degree'])
master_index = find_column_index(header_row, aliases=['master', 'masters', 'master degree'])
doctoral_index = find_column_index(header_row, aliases=['doctoral', 'doctorate', 'doctoral degree'])
workbook_type, workbook_type_label, workbook_type_reason = detect_workbook_type(header_row, {
    'id_number': id_number_index,
    'last_name': last_name_index,
    'first_name': first_name_index,
    'middle_name': middle_name_index,
    'department': department_index,
    'position': position_index,
    'date_hired': date_hired_index,
    'status': status_index,
    'email': email_index,
})
rows = []
for row in ws.iter_rows(min_row=header_row_index + 1, values_only=True):
    if not row:
        continue
    if all(value is None or clean(value) == '' for value in row[:11]):
        continue
    if is_header_like_row(row):
        continue
    last_name = clean(row[last_name_index] if last_name_index is not None and last_name_index < len(row) else (row[1] if len(row) > 1 else ''))
    first_name = clean(row[first_name_index] if first_name_index is not None and first_name_index < len(row) else (row[2] if len(row) > 2 else ''))
    middle_name = clean(row[middle_name_index] if middle_name_index is not None and middle_name_index < len(row) else (row[3] if len(row) > 3 else ''))
    if not last_name and not first_name:
        continue
    education = []
    for level, idx in [('Bachelor', bachelor_index), ('Master', master_index), ('Doctoral', doctoral_index)]:
        item = parse_education(level, row[idx] if idx is not None and idx < len(row) else None)
        if item:
            education.append(item)
    rows.append({
        'idNumber': clean(row[id_number_index] if id_number_index is not None and id_number_index < len(row) else ''),
        'importRowNumber': len(rows) + 1,
        'lastName': normalize_name(last_name),
        'firstName': normalize_name(first_name),
        'middleName': normalize_name(middle_name),
        'gender': normalize_gender(row[gender_index] if gender_index is not None and gender_index < len(row) else (row[4] if len(row) > 4 else '')),
        'department': normalize_department(row[department_index] if department_index is not None and department_index < len(row) else (row[5] if len(row) > 5 else '')),
        'position': strip_parenthetical(row[position_index] if position_index is not None and position_index < len(row) else (row[6] if len(row) > 6 else '')),
        'dateHired': normalize_date(row[date_hired_index] if date_hired_index is not None and date_hired_index < len(row) else (row[7] if len(row) > 7 else '')),
        'employmentStatus': normalize_status(row[status_index] if status_index is not None and status_index < len(row) else (row[9] if len(row) > 9 else '')),
        'email': clean(row[email_index] if email_index is not None and email_index < len(row) else (row[10] if len(row) > 10 else '')).lower(),
        'birthDate': normalize_birth_date(row[birth_date_index] if birth_date_index is not None and birth_date_index < len(row) else ''),
        'address': clean(row[address_index] if address_index is not None and address_index < len(row) else (row[14] if len(row) > 14 else '')),
        'contactNumber': normalize_contact_number(row[contact_number_index] if contact_number_index is not None and contact_number_index < len(row) else ''),
        'education': education,
        'documents': [],
        'staffCategory': 'Non-Teaching',
        'sourceClassification': classification,
        'sourceSheet': ws.title,
    })

print(json.dumps({
    'sheetName': ws.title,
    'classification': classification,
    'staffCategory': 'Non-Teaching',
    'workbookType': workbook_type,
    'workbookTypeLabel': workbook_type_label,
    'workbookTypeReason': workbook_type_reason,
    'records': rows,
}, ensure_ascii=False))
